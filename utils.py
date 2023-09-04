import openpyxl
# import mysql.connector
import pandas as pd
import datetime


def get_current_time():
    # 获取当前日期和时间
    current_datetime = datetime.datetime.now()
    # 格式化为YYMMDDHHMM格式
    formatted_datetime = current_datetime.strftime("%y%m%d%H%M")

    return formatted_datetime


class ReadFile():
    def __init__(self):
        pass

    def read_xlsx(self, file_name):

        return

    def read_xlsx_to_dict(self, file_path, sheet_name=None):
        data_dict_list = []

        # 打开XLSX文件
        workbook = openpyxl.load_workbook(file_path)

        # 选择工作表
        if sheet_name is None:
            sheet = workbook.active
        else:
            sheet = workbook[sheet_name]
        # 获取列头作为键
        header = [cell.value for cell in sheet[1]]

        # 遍历工作表中的行
        for row in sheet.iter_rows(min_row=2, values_only=True):
            data_dict = dict(zip(header, row))
            data_dict_list.append(data_dict)
        # 关闭工作簿
        workbook.close()
        return data_dict_list

    def read_xlsx_to_pandas(self, file_path, sheet_name=None):
        data = pd.read_excel(file_path, sheet_name=None)
        # if sheet_name is None:
        #     # 读取所有工作表
        #     data = pd.read_excel(file_path, sheet_name=None)
        # else:
        #     # 读取指定工作表
        #     data = pd.read_excel(file_path, sheet_name=sheet_name)
        return data


# class SaveData():
#     def __init__(self):
#         self.host = host
#         self.user = user
#         self.password = password
#         self.database = database
#         self.table = table
#     # 保存数据到文件
#     def save_to_file(self, data, file_path, file_format):
#         if file_format == 'xlsx':
#             data.to_excel(file_path, index=False)
#         elif file_format == 'json':
#             data.to_json(file_path, orient='records')

    # # 保存数据到数据库
    # def save_to_db(self, data):
    #     connection = mysql.connector.connect(
    #         host=self.host,
    #         user=self.user,
    #         password=self.password,
    #         database=self.database
    #     )
    #     cursor = connection.cursor()
    #
    #     for _, row in data.iterrows():
    #         values = tuple(row)
    #         placeholders = ', '.join(['%s'] * len(values))
    #         query = f"INSERT INTO {table} VALUES ({placeholders})"
    #         cursor.execute(query, values)
    #
    #     connection.commit()
    #     cursor.close()
    #     connection.close()

    # def save_data(self, data, save_to, file_path=None, table=None):
    #     if save_to == 'xlsx':
    #         data.to_excel(file_path, index=False)
    #     elif save_to == 'json':
    #         data.to_json(file_path, orient='records')
    #     elif save_to == 'csv':
    #         data.to_json(file_path, orient='records')
    #     elif save_to == 'db':
    #         connection = mysql.connector.connect(
    #             host=host,
    #             user=user,
    #             password=password,
    #             database=database
    #         )
    #         cursor = connection.cursor()
    #         # 假设你的数据是一个DataFrame，将数据逐行插入数据库表中
    #         for index, row in data.iterrows():
    #             values = tuple(row)
    #             query = f"INSERT INTO {table} VALUES {values}"
    #             cursor.execute(query)
    #         connection.commit()
    #         cursor.close()
    #         connection.close()


# class ExtractData():
#     # 从XLSX文件分块提取数据
#     def extract_data_from_xlsx(self, file_path, chunk_size=10000):
#         chunks = pd.read_excel(file_path, chunksize=chunk_size)
#         for chunk in chunks:
#             yield chunk
#
#     # 从MySQL数据库分块提取数据
#     def extract_data_from_mysql(self, host, user, password, database, query, chunk_size=10000):
#         connection = mysql.connector.connect(
#             host=host,
#             user=user,
#             password=password,
#             database=database
#         )
#         cursor = connection.cursor()
#         cursor.execute(query)
#
#         while True:
#             chunk = cursor.fetchmany(chunk_size)
#             if not chunk:
#                 break
#             yield pd.DataFrame(chunk, columns=cursor.column_names)
#
#         cursor.close()
#         connection.close()
